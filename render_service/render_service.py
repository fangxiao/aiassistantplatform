"""html_render 渲染 sidecar 服务(T18.17,独立容器承载)。

职责:HTML → PDF/PNG(headless Chromium/Playwright)+ PNG → PPTX(python-pptx)。
独立 sidecar 的动机:渲染引擎不进主 API 镜像(体积/依赖隔离),可独立扩缩。

契约(内网 HTTP,仅供平台 API 容器调用,不对外暴露端口):
    POST /render  {"html": str, "formats": ["pdf","png","pptx"], "scale": 2}
    → 200 {"files": {"pdf": "<b64>", "png": ["<b64>", ...], "pptx": "<b64>"}, "pages": int}
渲染约定(与插件侧 HTML 契约一致):
    - hash 路由 #print → 打印版单页(13.333x7.5in)PDF;
    - #/N → 第 N 页 PNG(Emulation.setDeviceMetricsOverride 控 scale)。
"""

from __future__ import annotations

import base64
import io
import os

from fastapi import FastAPI, HTTPException
from pydantic import BaseModel

PAGE_W, PAGE_H = 1280, 720  # 16:9 渲染视口(px),scale 控制实际分辨率

app = FastAPI(title="agentplatform-render-service")


class RenderRequest(BaseModel):
    html: str
    formats: list[str] = ["pdf"]
    scale: int = 2


def _data_url_html(html: str) -> str:
    """内联装载(绕过 CORS/网络):HTML 编码为 data URL 交给浏览器。"""
    return "data:text/html;charset=utf-8;base64," + base64.b64encode(html.encode("utf-8")).decode()


@app.get("/healthz")
async def healthz() -> dict:
    return {"ok": True}


@app.post("/render")
async def render(req: RenderRequest) -> dict:
    if not req.html or not req.html.strip():
        raise HTTPException(status_code=422, detail="html 为空")
    allowed = {"pdf", "png", "pptx"}
    formats = [f for f in req.formats if f in allowed]
    if not formats:
        raise HTTPException(status_code=422, detail=f"formats 需为 {sorted(allowed)} 的子集")

    from playwright.async_api import async_playwright

    files: dict[str, str | list[str]] = {}
    pages = 0
    async with async_playwright() as p:
        browser = await p.chromium.launch(args=["--no-sandbox"])
        try:
            page = await browser.new_page(
                viewport={"width": PAGE_W, "height": PAGE_H},
                device_scale_factor=max(1, min(req.scale, 4)),
            )
            await page.goto(_data_url_html(req.html), wait_until="networkidle")

            # 等待 hash 路由页数就绪(插件契约:body data-pages 或 #print)
            try:
                pages = int(await page.evaluate("document.body.dataset.pages || 0"))
            except Exception:  # noqa: BLE001
                pages = 0

            if "pdf" in formats:
                await page.goto(
                    _data_url_html(req.html + '<a id="__print_anchor"></a>') + "#print",
                    wait_until="networkidle",
                )
                pdf_bytes = await page.pdf(width="13.333in", height="7.5in", print_background=True)
                files["pdf"] = base64.b64encode(pdf_bytes).decode()

            if "png" in formats or "pptx" in formats:
                shots: list[str] = []
                total = pages or 1
                for n in range(1, total + 1):
                    await page.goto(_data_url_html(req.html) + f"#/{n}", wait_until="networkidle")
                    png = await page.screenshot(full_page=True)
                    shots.append(base64.b64encode(png).decode())
                files["png"] = shots
                pages = total

            if "pptx" in formats:
                from pptx import Presentation

                prs = Presentation()
                prs.slide_width = 12192000  # 16:9(EMU)
                prs.slide_height = 6858000
                blank = prs.slide_layouts[6]
                for b64 in files.get("png") or []:
                    slide = prs.slides.add_slide(blank)
                    slide.shapes.add_picture(io.BytesIO(base64.b64decode(b64)), 0, 0,
                                              width=prs.slide_width, height=prs.slide_height)
                buf = io.BytesIO()
                prs.save(buf)
                files["pptx"] = base64.b64encode(buf.getvalue()).decode()
        finally:
            await browser.close()

    return {"files": files, "pages": pages}


if __name__ == "__main":  # 本地直跑调试
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=int(os.environ.get("PORT", 8001)))
